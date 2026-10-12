import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY_DARK = RGBColor(15, 52, 96)       # #0F3460
    C_NAVY_DEEP = RGBColor(22, 33, 62)       # #16213E
    C_CORAL = RGBColor(233, 69, 96)          # #E94560
    C_PURPLE = RGBColor(83, 52, 131)         # #533483
    C_EMERALD = RGBColor(39, 174, 96)        # #27AE60
    C_BG_LIGHT = RGBColor(248, 249, 250)     # #F8F9FA
    C_CARD_BG = RGBColor(240, 243, 246)      # #F0F3F6
    C_WHITE = RGBColor(255, 255, 255)
    C_TEXT_DARK = RGBColor(26, 26, 46)       # #1A1A2E
    C_TEXT_MUTED = RGBColor(100, 110, 125)
    C_BORDER = RGBColor(220, 224, 230)

    def add_header(slide, title_text, category_text="FIRST REVIEW — OCTOBER 12, 2026"):
        # Top banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_CORAL
        top_bar.line.color.rgb = C_CORAL

        # Header box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_CORAL
        p_cat.space_after = Pt(2)

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_NAVY_DARK

        # Bottom subtle footer
        foot_bar = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
        tf_f = foot_bar.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
        p_foot = tf_f.paragraphs[0]
        p_foot.text = "Digital Logic Circuit Simulator | Team: Markanda Meher (206126020) & Nikhil (206126021) | Dept. of CSE"
        p_foot.font.size = Pt(9.5)
        p_foot.font.color.rgb = C_TEXT_MUTED

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.color.rgb = C_NAVY_DARK

    acc = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    acc.fill.solid()
    acc.fill.fore_color.rgb = C_CORAL
    acc.line.color.rgb = C_CORAL

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(10.9), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "ADVANCED DATA STRUCTURES & ALGORITHMS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Digital Logic Circuit Simulator"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Graph-Theoretic Circuit Simulation, Cycle Detection & Verification"
    p.font.size = Pt(17)
    p.font.color.rgb = RGBColor(210, 225, 245)
    p.space_after = Pt(28)

    # Divider bar
    p = tf1.add_paragraph()
    p.text = "━" * 48
    p.font.size = Pt(12)
    p.font.color.rgb = C_PURPLE
    p.space_after = Pt(24)

    p = tf1.add_paragraph()
    p.text = "Team Members:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CORAL

    p = tf1.add_paragraph()
    p.text = "• Markanda Meher  —  Roll No. 206126020\n• Nikhil                    —  Roll No. 206126021"
    p.font.size = Pt(14)
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(20)

    p = tf1.add_paragraph()
    p.text = "First Review — October 12, 2026"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(180, 195, 215)

    # =========================================================================
    # SLIDE 2: Agenda
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "First Review — Evaluation Structure & Agenda")

    agenda_items = [
        ("01", "Introduction & Problem Definition", "Limitations of physical breadboards, need for algorithmic software modeling."),
        ("02", "Comprehensive Literature Survey", "Critical analysis of 3 foundational papers (Kahn 1962, Tarjan 1972, Tarjan 1975) & drawbacks."),
        ("03", "System Block Diagram & Algorithmic Novelty", "Complete architectural pipeline, dual-direction graph representation & multi-pass execution."),
        ("04", "Algorithms & Detailed Pseudocode", "Pseudocode and walkthrough for DFS cycle detection, Kahn's topo-sort, Gate Eval, Truth Table."),
        ("05", "Standard Benchmark Dataset", "IEEE ISCAS '85 Suite (c17 benchmark netlist, .bench format, download links & complexity)."),
        ("06", "Implementation & Live Validation", "Working C++ simulator, static compilation, verification of Half Adder, MUX, and c17 benchmark.")
    ]

    for idx, (num, title, desc) in enumerate(agenda_items):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.4 + row * 1.75)
        
        create_card(s2, left, top, Inches(5.65), Inches(1.55))

        tb = s2.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.15), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{num}  |  {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Introduction - The Problem
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "1. Introduction — The Problem Definition")

    create_card(s3, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s3.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Physical Hardware Bottlenecks"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    bullets = [
        ("Component Fragility & Wire Clutter: ", "Physical breadboards suffer from loose jumper wires, contact bounce, and faulty DIP IC pins, making debugging circuit logic arduous."),
        ("Lack of Pre-Simulation Validation: ", "Accidental wiring of outputs to inputs creates unintended feedback loops, short circuits, and thermal runaway without warnings."),
        ("High Lab Cost & Limited Accessibility: ", "Oscilloscopes, logic analyzers, and discrete IC inventories represent a recurring financial barrier for academic labs."),
        ("Obscured Internal States: ", "Measuring signal states across 50+ intermediate nodes physically requires dozens of probes; software provides instantaneous introspection.")
    ]
    for b_title, b_text in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b_title
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + b_text
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(10)

    create_card(s3, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s3.shapes.add_textbox(Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "The Computer Science Problem"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(12)

    cs_bullets = [
        ("Graph-Theoretic Representation: ", "How can an electrical circuit with arbitrary interconnects and fan-out be modeled strictly using directed graphs?"),
        ("Asynchronous Evaluation Hazards: ", "In what mathematical sequence must gates be evaluated to guarantee every gate's inputs are ready before calculating its output?"),
        ("Structural Cycle Detection: ", "How can combinational loops be identified in O(V + E) time before simulation execution begins?"),
        ("Equivalence & Net Management: ", "How to group physically connected wire segments into single equipotential nets in near-constant O(alpha(N)) amortized time?")
    ]
    for b_title, b_text in cs_bullets:
        p = tf.add_paragraph()
        p.text = "• " + b_title
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + b_text
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 4: Introduction - Need to be Solved
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "1. Introduction — Need & Motivation")

    needs = [
        ("1. Pedagogy & Education", "Bridges the gap between Abstract Data Structures (Graphs, Disjoint Sets, Queues) and Real-World Computer Architecture.", "Students directly observe how Kahn's algorithm and DFS are applied in Electronic Design Automation (EDA)."),
        ("2. Algorithmic Pre-Validation", "Guarantees circuit correctness mathematically before signal propagation.", "Feedback loops and invalid fan-ins are caught instantly via DFS 3-colour marking, preventing infinite simulation loops."),
        ("3. High-Speed Combinatorial Verification", "Exhaustive truth table generation via bit-manipulation.", "Automates testing of all 2^N input combinations in microseconds, verifying functional equivalence against theoretical specs."),
        ("4. Zero-Dependency Standalone Architecture", "Engineered in modern C++17 with strict O(V + E) efficiency.", "No bloated GUI frameworks required; lightweight, cross-platform, and compiles to a fully self-contained static binary.")
    ]

    for idx, (title, sub, body) in enumerate(needs):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.4 + row * 2.7)
        create_card(s4, left, top, Inches(5.65), Inches(2.45))

        tb = s4.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), Inches(5.05), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 5: Literature Survey Overview
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "2. Literature Survey — Survey of 3 Research Papers")

    papers = [
        ("Paper 1 (1962)", "A. B. Kahn", "Communications of the ACM", "Topological Sorting of Large Networks", "Introduced in-degree queue-based topological sort for Directed Acyclic Graphs. Used in our simulator for dependency-ordered gate evaluation."),
        ("Paper 2 (1972)", "Robert E. Tarjan", "SIAM Journal on Computing", "Depth-First Search and Linear Graph Algorithms", "Established linear-time O(V+E) DFS graph traversal with node state classifications. Used in our simulator for 3-colour cycle detection."),
        ("Paper 3 (1975)", "Robert E. Tarjan", "Journal of the ACM (JACM)", "Efficiency of a Good but Not Linear Set Union Algorithm", "Proved O(m * alpha(n)) bound for Disjoint-Set Union with path compression & rank. Used in our simulator for wire-net management.")
    ]

    for idx, (tag, author, journal, p_title, summary) in enumerate(papers):
        top = Inches(1.4 + idx * 1.8)
        create_card(s5, Inches(0.8), top, Inches(11.733), Inches(1.6))

        tb = s5.shapes.add_textbox(Inches(1.05), top + Inches(0.15), Inches(11.2), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{tag}  |  {p_title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(3)

        p = tf.add_paragraph()
        p.text = f"Author: {author}  •  Publication: {journal}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = summary
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 6: Paper 1 Deep Dive & Drawbacks
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "2. Paper 1 Analysis — Kahn's Topological Sort (1962)")

    # Left: Details
    create_card(s6, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s6.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Paper Details & Core Algorithm"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(10)

    p = tf.add_paragraph()
    p.text = "Kahn, A. B. (1962). 'Topological sorting of large networks.' Communications of the ACM, 5(11), 558-562."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_PURPLE
    p.space_after = Pt(12)

    details1 = [
        "Core Mechanism: Computes in-degree for all vertices. Enqueues all vertices with in-degree = 0. Dequeues vertex, appends to sorted list, and decrements in-degree of downstream neighbours.",
        "Role in Simulator: Determines exact chronological calculation order so gate inputs are guaranteed valid before logic gate evaluation.",
        "Theoretical Bound: Runs in optimal O(V + E) time and O(V) auxiliary space."
    ]
    for d in details1:
        p = tf.add_paragraph()
        p.text = "• " + d
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

    # Right: Drawbacks
    create_card(s6, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s6.shapes.add_textbox(Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Critical Drawbacks & Limitations"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    drawbacks1 = [
        ("Drawback 1: Silent Failure on Cycles", "If the graph contains a cycle, Kahn's algorithm simply halts when the queue empties with processed count < V. It fails to identify WHICH specific gates or edges caused the cycle."),
        ("Drawback 2: Arbitrary Parallel Scheduling", "When multiple nodes have in-degree 0 simultaneously, FIFO queueing processes them arbitrarily, ignoring critical-path delays and timing constraints."),
        ("Drawback 3: Static Pre-computation Only", "If a single gate or input wire changes, the entire O(V + E) sort must be re-run from scratch (no incremental dynamic updating).")
    ]
    for title_d, body_d in drawbacks1:
        p = tf.add_paragraph()
        p.text = "⚠️ " + title_d
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = body_d
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 7: Paper 2 Deep Dive & Drawbacks
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "2. Paper 2 Analysis — Tarjan's DFS Framework (1972)")

    # Left: Details
    create_card(s7, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s7.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Paper Details & Core Algorithm"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(10)

    p = tf.add_paragraph()
    p.text = "Tarjan, R. E. (1972). 'Depth-First Search and Linear Graph Algorithms.' SIAM Journal on Computing, 1(2), 146-160."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_PURPLE
    p.space_after = Pt(12)

    details2 = [
        "Core Mechanism: Uses Depth-First Search with discovery timestamps to classify graph edges into tree edges, back edges, forward edges, and cross edges in linear O(V + E) time.",
        "Role in Simulator: We utilize a 3-colour marking scheme (White=Unvisited, Gray=Visiting, Black=Visited). A visit to a Gray node indicates a back edge, detecting combinational loops.",
        "Significance: Provided the mathematical foundation for linear-time directed graph analysis."
    ]
    for d in details2:
        p = tf.add_paragraph()
        p.text = "• " + d
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

    # Right: Drawbacks
    create_card(s7, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s7.shapes.add_textbox(Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Critical Drawbacks & Limitations"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    drawbacks2 = [
        ("Drawback 1: Call-Stack Overflow Risk", "Standard recursive DFS implementation consumes O(V) stack frames. Deeply cascaded circuits (thousands of logic levels) trigger OS stack overflow exceptions."),
        ("Drawback 2: Indiscriminate Cycle Rejection", "Tarjan's algorithm treats every back edge as an invalid cycle. In real sequential electronics, intentional cross-coupled feedback (SR latches, flip-flops) is required."),
        ("Drawback 3: Traversal Redundancy", "Running DFS purely for cycle detection requires traversing the entire graph, but does not yield the levelized evaluation order needed for zero-delay logic evaluation.")
    ]
    for title_d, body_d in drawbacks2:
        p = tf.add_paragraph()
        p.text = "⚠️ " + title_d
        p.font.bold = True
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = body_d
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 8: Paper 3 Deep Dive & Drawbacks
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "2. Paper 3 Analysis — Tarjan's Union-Find (1975)")

    # Left: Details
    create_card(s8, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s8.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Paper Details & Core Algorithm"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(10)

    p = tf.add_paragraph()
    p.text = "Tarjan, R. E. (1975). 'Efficiency of a Good but Not Linear Set Union Algorithm.' Journal of the ACM, 22(2), 215-225."
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_PURPLE
    p.space_after = Pt(12)

    details3 = [
        "Core Mechanism: Disjoint-Set forest with two key optimizations: Path Compression (flattens tree pointers during find) and Union by Rank (attaches shallow tree under taller tree root).",
        "Role in Simulator: Models physical wire networks ('nets'). When multiple wire segments tie together, Union-Find groups them into a single equipotential electrical net in O(alpha(N)) time.",
        "Significance: Proved the theoretical complexity bound involving the inverse Ackermann function alpha(N) <= 4."
    ]
    for d in details3:
        p = tf.add_paragraph()
        p.text = "• " + d
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

    # Right: Drawbacks
    create_card(s8, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s8.shapes.add_textbox(Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Critical Drawbacks & Limitations"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    drawbacks3 = [
        ("Drawback 1: Irreversible Merge Operations", "Classical Union-Find does NOT support dynamic splitting or disconnection ('un-unite'). If a user removes a wire, the disjoint sets must be fully rebuilt from scratch."),
        ("Drawback 2: Absence of Signal Directionality", "Union-Find models symmetric equivalence (undirected connectivity). It cannot distinguish between a signal DRIVER (output pin) and a signal SINK (input pin)."),
        ("Drawback 3: Memory Overhead for Small Circuits", "Maintaining parallel parent and rank vectors incurs unnecessary indirection when simple adjacency lists already track direct gate-to-gate connections.")
    ]
    for title_d, body_d in drawbacks3:
        p = tf.add_paragraph()
        p.text = "⚠️ " + title_d
        p.font.bold = True
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = body_d
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 9: Literature Survey Synthesis
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "2. Literature Survey — Synthesis & Research Gap Matrix")

    # Table of comparison
    table_shape = s9.shapes.add_table(5, 5, Inches(0.8), Inches(1.4), Inches(11.733), Inches(3.6))
    tbl = table_shape.table
    
    headers = ["Evaluation Metric", "Kahn (1962)", "Tarjan DFS (1972)", "Tarjan UF (1975)", "Our Integrated Simulator"]
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK if c < 4 else C_CORAL
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_WHITE

    rows_data = [
        ("Primary Capability", "Dependency Ordering", "Cycle / SCC Detection", "Net Equivalence", "End-to-End Simulation Pipeline"),
        ("Cycle Identification", "Fails without pin-pointing", "Pinpoints Back-Edge Nodes", "Cannot detect directed cycles", "Pre-validates & halts with exact node IDs"),
        ("Signal Directionality", "Directed edges only", "Directed traversal", "Undirected sets only", "Dual-Direction (Forward + Reverse edges)"),
        ("Simulation Suitability", "Ordering only (no logic)", "Traverse only (no eval)", "Connectivity only", "Complete Boolean evaluation & Truth Table")
    ]

    for r, rdata in enumerate(rows_data):
        for c, val in enumerate(rdata):
            cell = tbl.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if (r % 2 == 0) else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            if c == 4:
                p.font.bold = True
                p.font.color.rgb = C_NAVY_DARK

    # Bottom summary box
    create_card(s9, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.55), bg_color=RGBColor(232, 240, 254), border_color=C_NAVY_DARK)
    tb = s9.shapes.add_textbox(Inches(1.05), Inches(5.35), Inches(11.2), Inches(1.25))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "💡 Research Gap & Synthesis:"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(4)

    p = tf.add_paragraph()
    p.text = "None of the individual research papers offer an end-to-end simulation solution. Kahn's algorithm fails when cycles are present, Tarjan's DFS detects cycles but does not yield evaluation levels, and Union-Find ignores signal flow. Our contribution integrates all three algorithms into a robust, staged pipeline that validates, orders, and simulates circuits with optimal theoretical bounds."
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 10: Block Diagram
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "3. Block Diagram — End-to-End System Architecture")

    # 4 Architecture stages across the screen
    stages = [
        ("STAGE 1: GRAPH BUILDER", "Input Parsing & Topology", [
            "• Node allocation via addGate()",
            "• Wire routing via connect()",
            "• Forward Adjacency List adjList[u]",
            "• Reverse Dependency inputGateIds[v]",
            "• Union-Find wire net clustering"
        ], C_NAVY_DARK),
        ("STAGE 2: CYCLE VALIDATION", "DFS Pre-Verification", [
            "• 3-Colour Marking (W, G, B)",
            "• Back-edge detection (Gray visit)",
            "• Halts simulation if cycle detected",
            "• Prints diagnostic offending gate IDs",
            "• Guarantees valid DAG before execution"
        ], C_CORAL),
        ("STAGE 3: TOPOLOGICAL SCHEDULER", "Kahn's BFS Order", [
            "• In-degree calculation array",
            "• Zero-indegree queue initialization",
            "• BFS traversal with in-degree updates",
            "• Produces linear evaluation vector",
            "• Guarantees data ready at every gate"
        ], C_PURPLE),
        ("STAGE 4: EXECUTION ENGINE", "Boolean Eval & Truth Table", [
            "• Gate logic switch (AND, OR, NOT, etc.)",
            "• Multi-input bitwise reduction",
            "• Exhaustive 2^N combination loop",
            "• Bit-mask extraction ((combo>>k)&1)",
            "• Formatted truth table output generation"
        ], C_EMERALD)
    ]

    for idx, (title, sub, items, col_color) in enumerate(stages):
        left = Inches(0.8 + idx * 2.98)
        top = Inches(1.5)
        width = Inches(2.8)
        height = Inches(5.2)

        create_card(s10, left, top, width, height)

        # Header strip on card
        strip = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.65))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col_color
        strip.line.color.rgb = col_color

        tb = s10.shapes.add_textbox(left + Inches(0.1), top + Inches(0.08), width - Inches(0.2), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        # Subtitle and content
        tb_c = s10.shapes.add_textbox(left + Inches(0.15), top + Inches(0.75), width - Inches(0.3), Inches(4.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = sub
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(10)

        for it in items:
            p = tf_c.add_paragraph()
            p.text = it
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 11: Algorithmic Novelty & Contribution
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "3. Algorithmic Novelty & Technical Contribution")

    contribs = [
        ("1. Dual-Direction Graph Representation", "Maintains forward edges (adjList[u]) for traversal alongside reverse dependencies (gate.inputGateIds) simultaneously.", "Achieves O(1) gate insertion and O(1) wire connection, while allowing instantaneous fan-in retrieval during Boolean evaluation without expensive reverse lookups."),
        ("2. Staged Pre-Verification Pipeline", "Separates cycle detection from topological evaluation into distinct algorithmic phases.", "Standard tools crash or oscillate indefinitely when cycles exist. Our pre-simulation DFS phase strictly guarantees DAG properties, printing offending gate labels before simulation begins."),
        ("3. Combinatorial Exhaustion via Bit Manipulation", "Generates full truth tables across all 2^N input vectors using bitwise shifts ((combo >> bitPosition) & 1).", "Eliminates recursive truth-table overhead. Allows rapid verification of full truth tables (up to 32 rows for ISCAS c17) in under 1 millisecond."),
        ("4. Union-Find Disjoint Set Wire Management", "Integrates path compression and union by rank to manage multi-terminal wire nets.", "Abstracts multi-branch connections into unified electrical nodes with O(alpha(N)) amortized complexity, modeling physical PCB nets cleanly.")
    ]

    for idx, (title, sub, body) in enumerate(contribs):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.4 + row * 2.7)
        create_card(s11, left, top, Inches(5.65), Inches(2.45))

        tb = s11.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), Inches(5.05), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 12: Module 1 - Graph & Union-Find
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "4. Algorithms — Module 1: Graph & Union-Find")

    # Left: Pseudocode
    create_card(s12, Inches(0.8), Inches(1.4), Inches(6.0), Inches(5.35))
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.6), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithm 1: Graph Wire & Net Connection"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(8)

    code1 = """procedure ADD_GATE(type, label):
    id ← gates.size()
    gates.append(Gate(id, type, label))
    adjList.append([])  // empty forward edges
    return id

procedure CONNECT(fromId, toId):
    adjList[fromId].append(toId)        // Forward edge
    gates[toId].inputGateIds.append(fromId) // Reverse edge
    wireNets.unite(fromId, toId)       // Wire equivalence

procedure FIND(x):
    if parent[x] ≠ x then
        parent[x] ← FIND(parent[x])    // Path compression
    return parent[x]

procedure UNITE(x, y):
    rootX ← FIND(x); rootY ← FIND(y)
    if rootX == rootY then return false
    if rank[rootX] < rank[rootY] then parent[rootX] ← rootY
    else if rank[rootX] > rank[rootY] then parent[rootY] ← rootX
    else parent[rootY] ← rootX; rank[rootX] += 1
    return true"""

    p = tf.add_paragraph()
    p.text = code1
    p.font.size = Pt(9.5)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(30, 40, 55)

    # Right: Explanation
    create_card(s12, Inches(7.1), Inches(1.4), Inches(5.4), Inches(5.35))
    tb = s12.shapes.add_textbox(Inches(7.35), Inches(1.65), Inches(4.9), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithmic Walkthrough & Mechanics"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    walk_items = [
        ("Dual-Edge Maintenance: ", "Every wire insertion updates both adjList[from] and inputGateIds[to] in O(1) amortized time, providing bidirectional accessibility."),
        ("Path Compression: ", "During find(x), all traversed nodes repoint directly to the set representative, flattening the tree depth."),
        ("Union by Rank: ", "Guarantees maximum tree height increases only when two trees of identical rank merge, ensuring height <= log(N)."),
        ("Complexity: ", "Any sequence of M operations on N elements takes O(M * alpha(N)) time, which is practically linear.")
    ]
    for w_title, w_desc in walk_items:
        p = tf.add_paragraph()
        p.text = "• " + w_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + w_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 13: Module 2 - DFS Cycle Detection
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "4. Algorithms — Module 2: DFS 3-Colour Cycle Detection")

    # Left: Pseudocode
    create_card(s13, Inches(0.8), Inches(1.4), Inches(6.0), Inches(5.35))
    tb = s13.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.6), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithm 2: 3-Colour DFS Cycle Detection"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(8)

    code2 = """procedure HAS_CYCLE_DFS(node, color[], graph):
    color[node] ← GRAY // Currently on recursion stack

    for each neighbor in graph.adjList[node] do:
        if color[neighbor] == GRAY then
            print "Cycle detected: Back-edge to ", neighbor
            return true
        if color[neighbor] == WHITE then
            if HAS_CYCLE_DFS(neighbor, color, graph) then
                return true

    color[node] ← BLACK // Node fully processed
    return false

procedure HAS_CYCLE(graph):
    color ← array of size V initialized to WHITE
    for i ← 0 to V - 1 do:
        if color[i] == WHITE then
            if HAS_CYCLE_DFS(i, color, graph) then
                return true
    return false"""

    p = tf.add_paragraph()
    p.text = code2
    p.font.size = Pt(9.5)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(30, 40, 55)

    # Right: Explanation
    create_card(s13, Inches(7.1), Inches(1.4), Inches(5.4), Inches(5.35))
    tb = s13.shapes.add_textbox(Inches(7.35), Inches(1.65), Inches(4.9), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "3-Colour Node Classification"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    dfs_items = [
        ("WHITE (0) — Unvisited: ", "Gate has not yet been traversed by the DFS search algorithm."),
        ("GRAY (1) — Active in Stack: ", "Gate is currently being explored on the recursion call stack. Encounters with Gray nodes strictly define a BACK-EDGE (cycle)."),
        ("BLACK (2) — Fully Explored: ", "Gate and all downstream descendants have been validated as acyclic; skipped in subsequent traversals."),
        ("Disconnected Graph Handling: ", "Outer loop iterates across all V nodes, ensuring disconnected subcircuits are comprehensively tested."),
        ("Complexity: ", "Time Complexity: O(V + E) | Space Complexity: O(V).")
    ]
    for w_title, w_desc in dfs_items:
        p = tf.add_paragraph()
        p.text = "• " + w_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + w_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 14: Module 3 - Kahn's Topological Sort
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "4. Algorithms — Module 3: Kahn's Topological Sort")

    # Left: Pseudocode
    create_card(s14, Inches(0.8), Inches(1.4), Inches(6.0), Inches(5.35))
    tb = s14.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.6), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithm 3: Kahn's BFS Topological Sort"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(8)

    code3 = """procedure TOPOLOGICAL_SORT(graph):
    inDegree ← array of size V initialized to 0
    for i ← 0 to V - 1 do:
        for each neighbor in graph.adjList[i] do:
            inDegree[neighbor] ← inDegree[neighbor] + 1

    queue Q ← empty FIFO queue
    for i ← 0 to V - 1 do:
        if inDegree[i] == 0 then
            Q.push(i) // Seed with PRIMARY INPUT gates

    order ← empty vector
    while Q is not empty do:
        curr ← Q.front(); Q.pop()
        order.append(curr)

        for each neighbor in graph.adjList[curr] do:
            inDegree[neighbor] ← inDegree[neighbor] - 1
            if inDegree[neighbor] == 0 then
                Q.push(neighbor)

    return order"""

    p = tf.add_paragraph()
    p.text = code3
    p.font.size = Pt(9.5)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(30, 40, 55)

    # Right: Explanation
    create_card(s14, Inches(7.1), Inches(1.4), Inches(5.4), Inches(5.35))
    tb = s14.shapes.add_textbox(Inches(7.35), Inches(1.65), Inches(4.9), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Execution Mechanics & Guarantees"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(12)

    kahn_items = [
        ("Dependency Resolution: ", "In-degree corresponds to unresolved incoming wires. A gate enters the queue if and only if all its dependencies have been satisfied."),
        ("Evaluation Guarantee: ", "For every directed wire u -> v, gate u is guaranteed to precede gate v in the output array. No gate is evaluated with uncomputed inputs."),
        ("Double Cycle Check: ", "If order.size() < V, an un-ordered cyclic core exists; serves as a secondary validation safeguard."),
        ("Levelization: ", "BFS nature inherently levelizes circuit gates from Primary Inputs (Level 0) to Primary Outputs."),
        ("Complexity: ", "Time Complexity: O(V + E) | Space Complexity: O(V + E).")
    ]
    for w_title, w_desc in kahn_items:
        p = tf.add_paragraph()
        p.text = "• " + w_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + w_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 15: Module 4 & 5 - Gate Evaluation & Truth Table
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "4. Algorithms — Modules 4 & 5: Logic Eval & Truth Table")

    # Left: Evaluation Pseudocode
    create_card(s15, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s15.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.3), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithm 4: Gate Boolean Logic Evaluation"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(8)

    code4 = """procedure EVALUATE_GATE(gate, allGates):
    switch gate.type:
        case AND:
            res ← 1
            for inId in gate.inputGateIds do:
                res ← res & allGates[inId].outputValue
            gate.outputValue ← res
        case OR:
            res ← 0
            for inId in gate.inputGateIds do:
                res ← res | allGates[inId].outputValue
            gate.outputValue ← res
        case NOT:
            gate.outputValue ← !allGates[gate.inputGateIds[0]].val
        case XOR:
            res ← 0
            for inId in gate.inputGateIds do:
                res ← res ^ allGates[inId].outputValue
            gate.outputValue ← res
        case NAND:
            // Inverted AND logic
            gate.outputValue ← !(AND logic)
        case OUTPUT:
            gate.outputValue ← allGates[gate.inputGateIds[0]].val"""

    p = tf.add_paragraph()
    p.text = code4
    p.font.size = Pt(8.5)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(30, 40, 55)

    # Right: Truth Table Pseudocode
    create_card(s15, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s15.shapes.add_textbox(Inches(7.0), Inches(1.55), Inches(5.3), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Algorithm 5: Combinatorial Truth Table"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(8)

    code5 = """procedure GENERATE_TRUTH_TABLE(circuit):
    inputs ← filter(circuit.gates, type == INPUT)
    outputs ← filter(circuit.gates, type == OUTPUT)
    N ← inputs.size()
    totalRows ← 2^N

    for combo ← 0 to totalRows - 1 do:
        // Set Primary Inputs via bitmask
        for i ← 0 to N - 1 do:
            bitPos ← N - 1 - i
            inputs[i].outputValue ← (combo >> bitPos) & 1

        // Evaluate whole circuit (Topo order)
        SIMULATE(circuit, verbose = false)

        // Record output tuple
        printRow(inputs.values, outputs.values)

// Time Complexity: O(2^N * (V + E))
// Space Complexity: O(V)"""

    p = tf.add_paragraph()
    p.text = code5
    p.font.size = Pt(9.0)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(30, 40, 55)

    # =========================================================================
    # SLIDE 16: Section 5 - Benchmark Dataset
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "5. Dataset — IEEE ISCAS '85 Standard Logic Benchmarks")

    # Left: Explanation & Details
    create_card(s16, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s16.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "What is the ISCAS '85 Benchmark Suite?"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY_DARK
    p.space_after = Pt(10)

    dataset_bullets = [
        ("International Standard: ", "Published by IEEE International Symposium on Circuits and Systems (ISCAS); the worldwide benchmark for validating logic simulators, DAG algorithms, and ATPG tools."),
        ("Benchmark Format (.bench / .isc): ", "Netlist specification defining PRIMARY INPUTS, PRIMARY OUTPUTS, and Boolean gate assignments with explicit fan-ins."),
        ("Verification Methodology: ", "We feed official benchmark netlists into our simulator and compare computed truth tables against published IEEE reference tables."),
        ("Download Repository Links: ", "• University of Tennessee: https://pld.cs.utk.edu/bvls/iscas85/\n• Czech Technical University: https://ddd.fit.cvut.cz/prj/Benchmarks/\n• Trust-Hub Benchmark Suite: https://trust-hub.org/benchmarks")
    ]
    for b_title, b_text in dataset_bullets:
        p = tf.add_paragraph()
        p.text = "• " + b_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_CORAL
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        p2.text = "  " + b_text
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_DARK
        p2.space_after = Pt(8)

    # Right: Benchmark Circuit Table
    create_card(s16, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.35))
    tb = s16.shapes.add_textbox(Inches(7.05), Inches(1.65), Inches(5.2), Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ISCAS '85 Suite Circuit Specifications"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(10)

    table_shape_b = s16.shapes.add_table(6, 4, Inches(7.0), Inches(2.2), Inches(5.3), Inches(2.8))
    tbl_b = table_shape_b.table
    
    b_headers = ["Circuit", "Inputs", "Outputs", "Total Gates & Function"]
    for c, h in enumerate(b_headers):
        cell = tbl_b.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = C_WHITE

    bench_rows = [
        ("c17", "5", "2", "6 NAND gates (Included in Demo!)"),
        ("c432", "36", "7", "160 gates (Interrupt Controller)"),
        ("c499", "41", "32", "202 gates (32-bit Error Correction)"),
        ("c880", "60", "26", "383 gates (8-bit ALU)"),
        ("c1355", "41", "32", "546 gates (32-bit SEC Circuit)")
    ]
    for r, rdata in enumerate(bench_rows):
        for c, val in enumerate(rdata):
            cell = tbl_b.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if (r % 2 == 0) else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_DARK
            if r == 0:
                p.font.bold = True
                p.font.color.rgb = C_CORAL

    # Note below table
    tb_n = s16.shapes.add_textbox(Inches(7.0), Inches(5.3), Inches(5.3), Inches(1.2))
    tf_n = tb_n.text_frame
    tf_n.word_wrap = True
    p = tf_n.paragraphs[0]
    p.text = "⭐ Built-in Benchmark Demo: Our simulator includes a pre-configured implementation of the official ISCAS '85 c17 benchmark circuit (Option 11), generating all 32 truth table rows in < 1 ms."
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    # =========================================================================
    # SLIDE 17: Experimental Validation & Demos
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "Experimental Results — Verification & Benchmarks")

    demos = [
        ("Demo 1: Half Adder (6 Nodes)", "Sum = A XOR B\nCarry = A AND B", [
            "(0, 0) -> Sum=0, Carry=0",
            "(0, 1) -> Sum=1, Carry=0",
            "(1, 0) -> Sum=1, Carry=0",
            "(1, 1) -> Sum=0, Carry=1"
        ], C_NAVY_DARK),
        ("Demo 2: 2-to-1 MUX (8 Nodes)", "Y = (A & !S) | (B & S)\n3 Inputs (A, B, Sel)", [
            "Sel=0 -> Output Y = A",
            "Sel=1 -> Output Y = B",
            "Verified all 8 truth rows",
            "Zero timing anomalies"
        ], C_PURPLE),
        ("Demo 3: ISCAS c17 (13 Nodes)", "Standard IEEE Benchmark\n5 Inputs, 6 NAND, 2 Outputs", [
            "Verified 32/32 truth rows",
            "Zero combinational loops",
            "Simulation Time: < 1 ms",
            "100% IEEE spec match"
        ], C_EMERALD),
        ("Demo 4: Cycle Detection", "Back-Edge Error Detection\nGate 1 -> Gate 0 -> Gate 1", [
            "Traversed via DFS 3-colour",
            "Caught back-edge (Gray)",
            "Halts simulation safely",
            "Reports offending gate IDs"
        ], C_CORAL)
    ]

    for idx, (title, formula, points, col_color) in enumerate(demos):
        left = Inches(0.8 + idx * 2.98)
        top = Inches(1.5)
        width = Inches(2.8)
        height = Inches(5.2)

        create_card(s17, left, top, width, height)

        strip = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.65))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col_color
        strip.line.color.rgb = col_color

        tb = s17.shapes.add_textbox(left + Inches(0.1), top + Inches(0.08), width - Inches(0.2), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        tb_c = s17.shapes.add_textbox(left + Inches(0.15), top + Inches(0.75), width - Inches(0.3), Inches(4.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = formula
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(12)

        for pt in points:
            p = tf_c.add_paragraph()
            p.text = "✓ " + pt
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 18: Visual Console UI & Signal Flow
    # =========================================================================
    s17b = prs.slides.add_slide(blank_layout)
    add_header(s17b, "Visual Console Interface & Signal Flow Architecture")

    ui_features = [
        ("1. ASCII Circuit Schematics", "Visual circuit diagrams rendered directly in terminal for all preset circuits.", [
            "• Half Adder, NOR, 2-to-1 MUX & ISCAS c17",
            "• Dynamic flowcharts for custom circuits",
            "• Clear pin labels and directed wire paths",
            "• Zero GUI dependency, 100% cross-platform"
        ], C_NAVY_DARK),
        ("2. 3-Stage Signal Pipeline", "Clear topological separation of signal propagation stages.", [
            "• Stage 1: Primary Inputs (Data injected)",
            "• Stage 2: Logic Gates (Evaluations)",
            "• Stage 3: Primary Outputs (Probes)",
            "• Shows exact signal forwarding at each step"
        ], C_PURPLE),
        ("3. Step-by-Step Simulation", "Interactive execution trace showing gate arithmetic in real-time.", [
            "• Displays Kahn's topological sequence",
            "• Prints exact formula (e.g., 1 XOR 1 = 0)",
            "• Shows HIGH/LOW logic levels per wire",
            "• Optional interactive pause after each gate"
        ], C_EMERALD),
        ("4. Gate Inspector & Probe", "Deep diagnostic introspection of any chosen gate in O(1) time.", [
            "• Fan-In: Source gates & live signal values",
            "• Fan-Out: Downstream gates driven",
            "• Logic status (HIGH/LOW) & formula",
            "• Instantly identifies floating/dangling pins"
        ], C_CORAL)
    ]

    for idx, (title, sub, points, col_color) in enumerate(ui_features):
        left = Inches(0.8 + idx * 2.98)
        top = Inches(1.5)
        width = Inches(2.8)
        height = Inches(5.2)

        create_card(s17b, left, top, width, height)

        strip = s17b.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.65))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col_color
        strip.line.color.rgb = col_color

        tb = s17b.shapes.add_textbox(left + Inches(0.1), top + Inches(0.08), width - Inches(0.2), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

        tb_c = s17b.shapes.add_textbox(left + Inches(0.15), top + Inches(0.75), width - Inches(0.3), Inches(4.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = sub
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = col_color
        p.space_after = Pt(12)

        for pt in points:
            p = tf_c.add_paragraph()
            p.text = pt
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 19: Complexity Analysis
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    add_header(s18, "Theoretical Asymptotic Complexity Analysis")

    table_shape_c = s18.shapes.add_table(10, 4, Inches(0.8), Inches(1.4), Inches(11.733), Inches(4.6))
    tbl_c = table_shape_c.table

    c_headers = ["Simulator Operation", "Underlying Algorithm", "Time Complexity", "Auxiliary Space Complexity"]
    for c, h in enumerate(c_headers):
        cell = tbl_c.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_WHITE

    complexities = [
        ("Gate Insertion (addGate)", "Vector dynamic allocation", "O(1) amortized", "O(1)"),
        ("Wire Connection (connect)", "Adjacency list push_back", "O(1) amortized", "O(1)"),
        ("Cycle Detection (hasCycle)", "Tarjan's 3-Colour DFS", "O(V + E)", "O(V) color vector"),
        ("Evaluation Scheduling", "Kahn's BFS Topological Sort", "O(V + E)", "O(V + E) in-degrees & queue"),
        ("Gate Logic Evaluation", "Bitwise reduction over inputs", "O(fan-in)", "O(1)"),
        ("Full Circuit Simulation", "Topo-Sort + Gate Evaluation", "O(V + E)", "O(V)"),
        ("Truth Table Generation", "Exhaustive bitmask enumeration", "O(2^N * (V + E))", "O(V)"),
        ("Wire-Net Find (find)", "Path Compression Disjoint Set", "O(alpha(N)) ≈ O(1)", "O(1) call stack"),
        ("Wire-Net Union (unite)", "Union by Rank Disjoint Set", "O(alpha(N)) ≈ O(1)", "O(1)")
    ]

    for r, rdata in enumerate(complexities):
        for c, val in enumerate(rdata):
            cell = tbl_c.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if (r % 2 == 0) else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT_DARK
            if c >= 2:
                p.font.name = "Consolas"
                p.font.bold = True
                p.font.color.rgb = C_CORAL if c == 2 else C_NAVY_DARK

    p_note = s18.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.6))
    tf_n = p_note.text_frame
    p = tf_n.paragraphs[0]
    p.text = "Where V = Total Gates, E = Total Wire Connections, N = Number of Primary Inputs, and alpha(N) is the Inverse Ackermann function (alpha(N) <= 4 for all practical inputs)."
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 19: Work Plan & Milestones
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    add_header(s19, "Project Roadmap & Subsequent Milestones")

    phases = [
        ("First Review (Current)", "October 12, 2026", [
            "✓ Problem definition & literature survey",
            "✓ 3 foundational papers & drawback analysis",
            "✓ Core DAG data structure & dual-edge graph",
            "✓ DFS 3-colour cycle detection algorithm",
            "✓ Kahn's topological sort scheduler",
            "✓ ISCAS '85 c17 benchmark integration"
        ], C_EMERALD),
        ("Second Review", "November 2026", [
            "• ISCAS '85 .bench file parser implementation",
            "• Support for larger benchmarks (c432, c880)",
            "• Unit test suite & automated regression runner",
            "• Subcircuit hierarchical abstraction (macros)",
            "• Gate delay modeling (unit-delay model)",
            "• Performance profiling on 10,000+ nodes"
        ], C_CORAL),
        ("Final Review", "December 2026", [
            "• Web-based / Graphical GUI visualizer",
            "• Interactive wire placement and live probes",
            "• Export circuit netlists to JSON / Verilog",
            "• Sequential circuit support (flip-flops/latches)",
            "• Final project thesis & comprehensive report",
            "• Open-source repository release"
        ], C_PURPLE)
    ]

    for idx, (p_title, p_date, tasks, col_color) in enumerate(phases):
        left = Inches(0.8 + idx * 3.98)
        top = Inches(1.5)
        width = Inches(3.75)
        height = Inches(5.2)

        create_card(s19, left, top, width, height)

        strip = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.8))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col_color
        strip.line.color.rgb = col_color

        tb = s19.shapes.add_textbox(left + Inches(0.15), top + Inches(0.08), width - Inches(0.3), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p_d = tf.add_paragraph()
        p_d.text = p_date
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = RGBColor(220, 235, 255)

        tb_c = s19.shapes.add_textbox(left + Inches(0.2), top + Inches(0.95), width - Inches(0.4), Inches(4.1))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        for t in tasks:
            p = tf_c.add_paragraph()
            p.text = t
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(7)

    # =========================================================================
    # SLIDE 20: References
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    add_header(s20, "Academic & Technical References")

    refs = [
        ("1. Kahn, A. B. (1962). ", "'Topological sorting of large networks.' Communications of the ACM, Vol. 5, No. 11, pp. 558–562. DOI: 10.1145/368996.369025."),
        ("2. Tarjan, R. E. (1972). ", "'Depth-First Search and Linear Graph Algorithms.' SIAM Journal on Computing, Vol. 1, No. 2, pp. 146–160. DOI: 10.1137/0201010."),
        ("3. Tarjan, R. E. (1975). ", "'Efficiency of a Good but Not Linear Set Union Algorithm.' Journal of the ACM (JACM), Vol. 22, No. 2, pp. 215–225. DOI: 10.1145/321879.321884."),
        ("4. Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2009). ", "Introduction to Algorithms (CLRS), 3rd Edition. MIT Press & McGraw-Hill. (Chapters 21 & 22: Graph Algorithms & Disjoint Sets)."),
        ("5. Mano, M. M., & Ciletti, M. D. (2013). ", "Digital Design: With an Introduction to the Verilog HDL, 5th Edition. Pearson Education."),
        ("6. Brglez, F., & Fujiwara, H. (1985). ", "'A neutral netlist of 10 combinational benchmark circuits and a target translator in FORTRAN.' Special Session on ATPG and Fault Simulation, IEEE ISCAS'85."),
        ("7. Burch, C. (2002). ", "'Logisim: A graphical system for logic circuit design and simulation.' Journal of Computing Sciences in Colleges, Vol. 17, No. 4, pp. 42–48.")
    ]

    create_card(s20, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.35))
    tb_r = s20.shapes.add_textbox(Inches(1.05), Inches(1.6), Inches(11.2), Inches(4.9))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    for r_num, r_body in refs:
        p = tf_r.add_paragraph()
        p.text = r_num
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK
        
        # Add remainder
        run = p.add_run()
        run.text = r_body
        run.font.bold = False
        run.font.size = Pt(10.5)
        run.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 21: Conclusion & Q&A
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    bg21 = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg21.fill.solid()
    bg21.fill.fore_color.rgb = C_NAVY_DARK
    bg21.line.color.rgb = C_NAVY_DARK

    acc21 = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    acc21.fill.solid()
    acc21.fill.fore_color.rgb = C_CORAL
    acc21.line.color.rgb = C_CORAL

    tb21 = s21.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    tf21 = tb21.text_frame
    tf21.word_wrap = True

    p = tf21.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(14)

    p = tf21.add_paragraph()
    p.text = "Digital Logic Circuit Simulator — Review 1 Complete"
    p.font.size = Pt(20)
    p.font.color.rgb = C_CORAL
    p.space_after = Pt(24)

    p = tf21.add_paragraph()
    p.text = "Questions & Discussion"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(210, 225, 245)
    p.space_after = Pt(24)

    p = tf21.add_paragraph()
    p.text = "Markanda Meher (206126020)  •  Nikhil (206126021)\nDepartment of Computer Science & Engineering"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(180, 195, 215)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Projects\DigitalLogicSimulator"
    out_file = os.path.join(out_dir, "Digital_Logic_Simulator_First_Review.pptx")
    create_deck(out_file)
